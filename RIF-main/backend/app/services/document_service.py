"""Phase 8 local document upload, native extraction, and PaddleOCR ingestion service."""

from __future__ import annotations

import hashlib
import io
import re
import shutil
import tempfile
import time
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from uuid import UUID

from sqlalchemy import desc, select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.constants import ResponseStatus, UserRole
from app.models.operations import ActionLog, Document, DocumentIngestionJob
from app.ocr.paddle_ocr import PaddleOcrRuntimeError, paddle_ocr_adapter


class DocumentServiceError(Exception):
    """Expected document-processing problem with a frontend-safe response."""

    def __init__(self, *, status: ResponseStatus, code: str, message: str, details: list[dict[str, Any]] | None = None) -> None:
        self.status = status
        self.code = code
        self.message = message
        self.details = details or []
        super().__init__(message)


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    page_count: int | None
    ocr_used: bool
    extraction_method: str
    ocr_duration_ms: int | None


@dataclass(frozen=True)
class DocumentIngestionResult:
    document: Document
    job: DocumentIngestionJob
    extraction: ExtractedDocument
    action_log_id: int | None


class DocumentService:
    """Store allowed local files, extract text, and record ingestion metadata."""

    # JPG and JPEG are stored under one canonical application type.
    ALLOWED_TYPES = {"pdf", "docx", "txt", "jpg", "png"}
    ALLOWED_EXTENSIONS = {
        ".pdf": "pdf",
        ".docx": "docx",
        ".txt": "txt",
        ".jpg": "jpg",
        ".jpeg": "jpg",
        ".png": "png",
    }
    CONTENT_TYPE_HINTS = {
        "pdf": {"application/pdf"},
        "docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/zip"},
        "txt": {"text/plain"},
        "jpg": {"image/jpeg", "image/jpg"},
        "png": {"image/png"},
    }

    def __init__(self) -> None:
        self.settings = get_settings()

    @property
    def upload_root(self) -> Path:
        root = self.settings.resolved_upload_directory / "documents"
        root.mkdir(parents=True, exist_ok=True)
        return root

    @property
    def text_root(self) -> Path:
        root = self.settings.resolved_upload_directory / "extracted_text"
        root.mkdir(parents=True, exist_ok=True)
        return root

    @staticmethod
    def _safe_filename(filename: str) -> str:
        base = Path(filename or "upload").name
        base = re.sub(r"[^A-Za-z0-9._ -]", "_", base).strip(" .")
        if not base:
            raise DocumentServiceError(
                status=ResponseStatus.VALIDATION_FAILED,
                code="invalid_filename",
                message="The uploaded filename is not valid.",
            )
        return base[:180]

    @classmethod
    def _canonical_type(cls, filename: str) -> str:
        extension = Path(filename).suffix.lower()
        document_type = cls.ALLOWED_EXTENSIONS.get(extension)
        if document_type is None:
            allowed = ", ".join(sorted(cls.ALLOWED_TYPES))
            raise DocumentServiceError(
                status=ResponseStatus.VALIDATION_FAILED,
                code="unsupported_file_type",
                message=f"Unsupported file type. Allowed types: {allowed}.",
                details=[{"filename": filename, "extension": extension or None}],
            )
        return document_type

    @staticmethod
    def _is_docx_bytes(content: bytes) -> bool:
        if not content.startswith(b"PK"):
            return False
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                names = set(archive.namelist())
            return "[Content_Types].xml" in names and "word/document.xml" in names
        except zipfile.BadZipFile:
            return False

    @classmethod
    def _validate_signature(cls, document_type: str, content: bytes) -> None:
        valid = {
            "pdf": content.startswith(b"%PDF-"),
            "png": content.startswith(b"\x89PNG\r\n\x1a\n"),
            "jpg": content.startswith(b"\xff\xd8\xff"),
            "docx": cls._is_docx_bytes(content),
            "txt": b"\x00" not in content,
        }.get(document_type, False)
        if not valid:
            raise DocumentServiceError(
                status=ResponseStatus.VALIDATION_FAILED,
                code="file_signature_mismatch",
                message="The file contents do not match the selected file type.",
                details=[{"file_type": document_type}],
            )

    @classmethod
    def validate_upload(cls, *, filename: str, content_type: str | None, content: bytes) -> tuple[str, str]:
        """Validate extension, content-type hint, and binary signature before saving."""

        safe_filename = cls._safe_filename(filename)
        document_type = cls._canonical_type(safe_filename)
        if not content:
            raise DocumentServiceError(
                status=ResponseStatus.VALIDATION_FAILED,
                code="empty_upload",
                message="The uploaded file is empty.",
            )
        if content_type and content_type.lower() not in {"application/octet-stream", "binary/octet-stream"}:
            allowed = cls.CONTENT_TYPE_HINTS[document_type]
            if content_type.lower() not in allowed:
                raise DocumentServiceError(
                    status=ResponseStatus.VALIDATION_FAILED,
                    code="content_type_mismatch",
                    message="The uploaded MIME type does not match the file extension.",
                    details=[{"declared_content_type": content_type, "expected_type": document_type}],
                )
        cls._validate_signature(document_type, content)
        return safe_filename, document_type

    @staticmethod
    def _timestamp() -> datetime:
        return datetime.now(timezone.utc)

    @staticmethod
    def _as_summary(document: Document) -> dict[str, Any]:
        return {
            "document_id": document.id,
            "original_filename": document.original_filename,
            "file_type": document.file_type,
            "file_size_bytes": document.file_size_bytes,
            "page_count": document.page_count,
            "ocr_used": document.ocr_used,
            "ingestion_status": document.ingestion_status,
            "extracted_text_available": bool(document.extracted_text_path),
            "document_version": document.document_version,
            "error_message": document.error_message,
            "created_at": document.created_at.isoformat(),
            "updated_at": document.updated_at.isoformat(),
        }

    def _write_binary(self, *, document_id: UUID, filename: str, content: bytes) -> Path:
        target_dir = self.upload_root / str(document_id)
        target_dir.mkdir(parents=True, exist_ok=True)
        target_path = target_dir / filename
        target_path.write_bytes(content)
        return target_path

    def _write_extracted_text(self, *, document_id: UUID, text: str) -> Path:
        target_path = self.text_root / f"{document_id}.txt"
        target_path.write_text(text, encoding="utf-8")
        return target_path

    @staticmethod
    def _extract_txt(path: Path) -> ExtractedDocument:
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
        text = text.strip()
        if not text:
            raise DocumentServiceError(
                status=ResponseStatus.OCR_FAILED,
                code="no_text_extracted",
                message="No readable text was found in the TXT file.",
            )
        return ExtractedDocument(text=text, page_count=1, ocr_used=False, extraction_method="native_txt", ocr_duration_ms=None)

    @staticmethod
    def _extract_docx(path: Path) -> ExtractedDocument:
        try:
            from docx import Document as WordDocument
        except ImportError as exc:  # pragma: no cover - dependency issue only
            raise DocumentServiceError(
                status=ResponseStatus.TOOL_FAILED,
                code="docx_library_unavailable",
                message="DOCX support is not installed. Install the Phase 8 requirements, then retry.",
            ) from exc
        try:
            word_document = WordDocument(str(path))
            parts: list[str] = [paragraph.text.strip() for paragraph in word_document.paragraphs if paragraph.text.strip()]
            for table_index, table in enumerate(word_document.tables, start=1):
                rows = []
                for row in table.rows:
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    rows.append(" | ".join(cells))
                if rows:
                    parts.append(f"[TABLE {table_index}]\n" + "\n".join(rows))
            text = "\n\n".join(parts).strip()
        except Exception as exc:
            raise DocumentServiceError(
                status=ResponseStatus.TOOL_FAILED,
                code="docx_extraction_failed",
                message="The DOCX file could not be read.",
            ) from exc
        if not text:
            raise DocumentServiceError(
                status=ResponseStatus.OCR_FAILED,
                code="no_text_extracted",
                message="No readable text was found in the DOCX file.",
            )
        return ExtractedDocument(text=text, page_count=None, ocr_used=False, extraction_method="native_docx", ocr_duration_ms=None)

    def _ocr_image(self, image_path: Path) -> tuple[str, int]:
        started = time.perf_counter()
        try:
            text = paddle_ocr_adapter.extract_text(image_path)
        except PaddleOcrRuntimeError as exc:
            raise DocumentServiceError(
                status=ResponseStatus.OCR_FAILED,
                code=exc.code,
                message=exc.message,
            ) from exc
        elapsed_ms = int((time.perf_counter() - started) * 1000)
        return text, elapsed_ms

    def _extract_image(self, path: Path) -> ExtractedDocument:
        text, elapsed_ms = self._ocr_image(path)
        return ExtractedDocument(text=text, page_count=1, ocr_used=True, extraction_method="paddleocr_image", ocr_duration_ms=elapsed_ms)

    def _extract_pdf(self, path: Path) -> ExtractedDocument:
        try:
            import fitz  # PyMuPDF
        except ImportError as exc:  # pragma: no cover - dependency issue only
            raise DocumentServiceError(
                status=ResponseStatus.TOOL_FAILED,
                code="pymupdf_unavailable",
                message="PDF support is not installed. Install the Phase 8 requirements, then retry.",
            ) from exc
        try:
            pdf = fitz.open(str(path))
        except Exception as exc:
            raise DocumentServiceError(
                status=ResponseStatus.TOOL_FAILED,
                code="pdf_open_failed",
                message="The PDF file could not be opened.",
            ) from exc

        page_sections: list[str] = []
        ocr_used = False
        total_ocr_ms = 0
        try:
            with tempfile.TemporaryDirectory(prefix="b2b_phase8_pdf_") as temp_dir:
                temp_path = Path(temp_dir)
                for page_number, page in enumerate(pdf, start=1):
                    native_text = page.get_text("text", sort=True).strip()
                    page_text = native_text
                    if len(native_text) < self.settings.document_native_text_min_characters:
                        image_path = temp_path / f"page-{page_number}.png"
                        pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                        pixmap.save(str(image_path))
                        page_text, elapsed_ms = self._ocr_image(image_path)
                        ocr_used = True
                        total_ocr_ms += elapsed_ms
                    if page_text.strip():
                        page_sections.append(f"[PAGE {page_number}]\n{page_text.strip()}")
        finally:
            pdf.close()

        text = "\n\n".join(page_sections).strip()
        if not text:
            raise DocumentServiceError(
                status=ResponseStatus.OCR_FAILED,
                code="no_text_extracted",
                message="No readable text was found in the PDF file.",
            )
        return ExtractedDocument(
            text=text,
            page_count=len(page_sections) if page_sections else None,
            ocr_used=ocr_used,
            extraction_method="paddleocr_scanned_pdf" if ocr_used else "native_pdf",
            ocr_duration_ms=total_ocr_ms if ocr_used else None,
        )

    def extract_document(self, *, path: Path, document_type: str) -> ExtractedDocument:
        if document_type == "txt":
            return self._extract_txt(path)
        if document_type == "docx":
            return self._extract_docx(path)
        if document_type == "pdf":
            return self._extract_pdf(path)
        if document_type in {"jpg", "png"}:
            return self._extract_image(path)
        raise DocumentServiceError(
            status=ResponseStatus.VALIDATION_FAILED,
            code="unsupported_file_type",
            message="This file type is not supported for document ingestion.",
        )

    def _write_action_log(
        self,
        db: Session,
        *,
        request_id: str,
        actor_role: UserRole,
        document: Document,
        status: str,
        error_message: str | None = None,
    ) -> ActionLog:
        log = ActionLog(
            request_id=request_id,
            session_id=None,
            pending_action_id=None,
            actor_role=actor_role.value,
            action_type="document_upload",
            target_table="documents",
            affected_record_ids={"document_id": str(document.id)},
            generated_sql=None,
            confirmation_status="not_required",
            status=status,
            error_message=error_message,
            created_at=self._timestamp(),
        )
        db.add(log)
        db.flush()
        return log

    def ingest_upload(
        self,
        db: Session,
        *,
        filename: str,
        content_type: str | None,
        content: bytes,
        request_id: str,
        actor_role: UserRole,
    ) -> DocumentIngestionResult:
        if len(content) > self.settings.document_max_upload_bytes:
            raise DocumentServiceError(
                status=ResponseStatus.VALIDATION_FAILED,
                code="file_too_large",
                message=f"The uploaded file exceeds the configured {self.settings.document_max_upload_bytes} byte limit.",
            )
        safe_filename, document_type = self.validate_upload(filename=filename, content_type=content_type, content=content)
        file_hash = hashlib.sha256(content).hexdigest()
        existing = db.scalar(select(Document).where(Document.file_hash == file_hash))
        if existing is not None:
            raise DocumentServiceError(
                status=ResponseStatus.DUPLICATE_DETECTED,
                code="duplicate_document",
                message="This exact document has already been uploaded.",
                details=[{"existing_document_id": str(existing.id), "original_filename": existing.original_filename}],
            )

        document = Document(
            original_filename=safe_filename,
            file_hash=file_hash,
            file_type=document_type,
            file_size_bytes=len(content),
            page_count=None,
            ocr_used=False,
            ingestion_status="processing",
            extracted_text_path=None,
            error_message=None,
            document_version=1,
        )
        db.add(document)
        db.flush()
        job = DocumentIngestionJob(
            document_id=document.id,
            status="processing",
            started_at=self._timestamp(),
            completed_at=None,
            ocr_duration_ms=None,
            error_message=None,
            created_at=self._timestamp(),
        )
        db.add(job)
        db.flush()

        try:
            binary_path = self._write_binary(document_id=document.id, filename=safe_filename, content=content)
            extraction = self.extract_document(path=binary_path, document_type=document_type)
            text_path = self._write_extracted_text(document_id=document.id, text=extraction.text)
            document.page_count = extraction.page_count
            document.ocr_used = extraction.ocr_used
            document.ingestion_status = "completed"
            document.extracted_text_path = str(text_path.relative_to(self.settings.resolved_upload_directory))
            document.error_message = None
            job.status = "completed"
            job.completed_at = self._timestamp()
            job.ocr_duration_ms = extraction.ocr_duration_ms
            job.error_message = None
            action_log = self._write_action_log(
                db,
                request_id=request_id,
                actor_role=actor_role,
                document=document,
                status="completed",
            )
            db.commit()
            db.refresh(document)
            db.refresh(job)
            return DocumentIngestionResult(document=document, job=job, extraction=extraction, action_log_id=action_log.id)
        except DocumentServiceError as exc:
            document.ingestion_status = "failed"
            document.error_message = exc.message
            job.status = "failed"
            job.completed_at = self._timestamp()
            job.error_message = exc.message
            self._write_action_log(
                db,
                request_id=request_id,
                actor_role=actor_role,
                document=document,
                status="failed",
                error_message=exc.message,
            )
            db.commit()
            raise
        except Exception as exc:  # pragma: no cover - defensive persistence error
            document.ingestion_status = "failed"
            document.error_message = "Document ingestion failed unexpectedly."
            job.status = "failed"
            job.completed_at = self._timestamp()
            job.error_message = "Document ingestion failed unexpectedly."
            self._write_action_log(
                db,
                request_id=request_id,
                actor_role=actor_role,
                document=document,
                status="failed",
                error_message=str(exc),
            )
            db.commit()
            raise DocumentServiceError(
                status=ResponseStatus.TOOL_FAILED,
                code="document_ingestion_failed",
                message="The document could not be ingested. Check the local log using the request ID.",
            ) from exc

    def list_documents(self, db: Session, *, limit: int = 50) -> list[dict[str, Any]]:
        documents = list(db.scalars(select(Document).order_by(desc(Document.created_at)).limit(limit)).all())
        return [self._as_summary(document) for document in documents]

    def get_document(self, db: Session, document_id: UUID) -> dict[str, Any] | None:
        document = db.get(Document, document_id)
        return self._as_summary(document) if document is not None else None

    def get_extracted_text(self, db: Session, document_id: UUID, *, max_characters: int = 20_000) -> dict[str, Any] | None:
        document = db.get(Document, document_id)
        if document is None:
            return None
        if not document.extracted_text_path:
            return {
                "document": self._as_summary(document),
                "text": None,
                "truncated": False,
                "available": False,
            }
        path = self.settings.resolved_upload_directory / document.extracted_text_path
        if not path.exists():
            return {
                "document": self._as_summary(document),
                "text": None,
                "truncated": False,
                "available": False,
            }
        text = path.read_text(encoding="utf-8")
        return {
            "document": self._as_summary(document),
            "text": text[:max_characters],
            "truncated": len(text) > max_characters,
            "available": True,
            "character_count": len(text),
        }


document_service = DocumentService()
